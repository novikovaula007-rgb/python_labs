import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.lab05.json_csv import check_extension


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX (через openpyxl).

    Первая строка CSV - заголовок. Лист называется "Sheet1".
    Ширина каждой колонки - по самой длинной строке в колонке,
    но не меньше 8 символов.
    Все значения записываются как текст (так же, как они лежат в CSV).

    Raises:
        неверное расширение файла -> ValueError
        CSV пустой или без строк с данными -> ValueError
        файл не найден -> FileNotFoundError
    """
    check_extension(csv_path, ".csv")
    check_extension(xlsx_path, ".xlsx")

    rows = []
    with open(csv_path, encoding="utf-8", newline="") as f:
        for row in csv.reader(f):
            if row:
                rows.append(row)

    # нужна хотя бы строка заголовка и одна строка данных
    if len(rows) < 2:
        raise ValueError("CSV пустой или в нём нет данных (только заголовок)")

    widths = {}
    for row in rows:
        for index, cell in enumerate(row):
            if len(cell) > widths.get(index, 0):
                widths[index] = len(cell)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for index, width in widths.items():
        letter = get_column_letter(index + 1)
        ws.column_dimensions[letter].width = max(8, width)

    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)

    wb.save(xlsx_path)


if __name__ == "__main__":
    csv_to_xlsx("data/lab05/samples/people.csv", "data/lab05/out/people.xlsx")
    csv_to_xlsx("data/lab05/samples/cities.csv", "data/lab05/out/cities.xlsx")
    print("XLSX-файлы созданы.")